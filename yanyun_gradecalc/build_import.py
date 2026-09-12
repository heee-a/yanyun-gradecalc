"""流派导入器：从毕业率计算器 Excel 提取流派关注词条，生成 builds/<流派>.json。

计算器的「期望」表 A 列输入区（第 2~24 行）列出了该流派建模用到的全部面板属性，
即该流派关心的词条集合。新增流派计算器后重新运行 import-build 即可。
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

from .stats_dict import _PANEL_TO_AFFIX

# 输入区下边界：A 列出现这些标签之后不再是属性输入
_STOP_LABELS = {"食物加成", "第三心法", "笛", "机制"}
_INPUT_MIN_ROW, _INPUT_MAX_ROW = 2, 24


def _build_name(filename: str) -> str:
    """从文件名提取流派名：'破竹樽110阶竞速轴属性毕业率进阶计算器2.0.xlsx' -> 破竹樽"""
    stem = re.sub(r"[.\d]+$", "", Path(filename).stem)  # 去掉尾部版本号（如 2.3）
    m = re.match(r"^([\u4e00-\u9fa5]{2,6}?)\d*阶", stem)
    return m.group(1) if m else stem.split("110")[0]


def extract_relevant_stats(xlsx_path: str | Path) -> dict:
    """提取计算器输入区的面板属性 -> 归纳为装备词条集合。

    属性标签有两个来源：第 1 行 B:E 的列头（最小攻击/最大攻击/穿透/伤害加成，
    对应第 2 行外功等主属性行的数值），以及 A 列第 2~24 行的行标签（会心率、
    首领增、各武器增伤、特技定音等）。
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = wb["期望"] if "期望" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(min_row=1, max_row=_INPUT_MAX_ROW, min_col=1, max_col=5))
    wb.close()

    panel_stats: list[str] = []
    for cell in rows[0][1:]:  # 第 1 行列头
        if isinstance(cell.value, str) and cell.value.strip() in _PANEL_TO_AFFIX:
            panel_stats.append(cell.value.strip())
    for row in rows[1:]:
        label = row[0].value
        if label is None:
            continue
        label = str(label).strip()
        if label in _STOP_LABELS:
            break
        has_value = any(isinstance(c.value, (int, float)) for c in row[1:])
        # 特技定音类（如 蓄力技定音/酩酊技定音）即使无 B 值也保留
        if label in _PANEL_TO_AFFIX or has_value or label.endswith("定音"):
            panel_stats.append(label)

    affixes: list[str] = []
    for p in panel_stats:
        for affix in _PANEL_TO_AFFIX.get(p, []):
            if affix not in affixes:
                affixes.append(affix)
        # 蓄力技定音 -> 蓄力技增伤 类特技
        if p.endswith("定音"):
            special = p[:-2] + "增伤"
            if special not in affixes:
                affixes.append(special)
    return {"panel_stats": panel_stats, "affixes": affixes}


def import_build(xlsx_path: str | Path, out_dir: str | Path = "builds") -> Path:
    """导入单个计算器，生成 builds/<流派>.json，返回输出路径。"""
    xlsx_path = Path(xlsx_path)
    info = extract_relevant_stats(xlsx_path)
    name = _build_name(xlsx_path.name)
    weights = {a: 1.0 for a in info["affixes"]}
    data = {
        "name": name,
        "source_file": xlsx_path.name,
        "note": "权重默认均为 1.0，可按流派重要性自行调整；权重 0 的词条不参与毕业度计算",
        "panel_stats": info["panel_stats"],
        "affix_weights": weights,
    }
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    path = out / f"{name}.json"
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def import_folder(path: str | Path, out_dir: str | Path = "builds") -> list[Path]:
    """批量导入目录（含子目录）下所有 xlsx 计算器。"""
    path = Path(path)
    files = sorted(path.rglob("*.xlsx")) if path.is_dir() else [path]
    if not files:
        raise FileNotFoundError(f"{path} 下没有找到 xlsx 文件")
    return [import_build(f, out_dir) for f in files
            if not f.name.startswith("~$")]
