"""yanyun-gradecalc 单元测试（不依赖 OCR 模型）。pytest -q"""

from __future__ import annotations

import json

import pytest
from openpyxl import Workbook

from yanyun_gradecalc.build_import import _build_name, extract_relevant_stats, import_build
from yanyun_gradecalc.ocr import parse_ocr_result
from yanyun_gradecalc.scoring import (affix_roll, compare_builds, grade_of,
                                      load_max_table, score_piece)
from yanyun_gradecalc.stats_dict import normalize_stat, parse_value, split_affix_marker


def box(x, y, text, h=22):
    """合成 OCR 框：左上(x,y) 右下(x+w,y+h)。"""
    return [[x, y], [x + 30 * len(text), y], [x + 30 * len(text), y + h], [x, y + h]], text, 0.9


# ---------------- stats_dict ----------------
def test_normalize_stat():
    assert normalize_stat("会心率") == "会心率"
    assert normalize_stat("最 大 外 功 攻 击") == "最大外功攻击"
    assert normalize_stat("无名剑法·蓄力技增伤") == "无名剑法·蓄力技增伤"
    assert normalize_stat("首领增伤") == "对首领单位增伤"
    assert normalize_stat("乱七八糟词条") is None


def test_split_marker_and_value():
    name, conv, rec = split_affix_marker("转会心率荐")
    assert name == "会心率" and conv and rec
    assert parse_value("12.7%") == (12.7, "percent")
    assert parse_value("114.1") == (114.1, "flat")
    assert parse_value("外功") is None


# ---------------- ocr parse ----------------
def test_parse_angled_layout_pairs_correctly():
    """模拟斜拍照片：数值列在右侧 x≈900，左右列存在系统性纵向偏移。"""
    ocr = [
        box(300, 100, "造谐"), box(370, 220, "1476"),
        box(310, 320, "装备等阶"),
        box(310, 410, "气血最大值"), box(950, 580, "8750"),
        box(280, 470, "外功防御"), box(950, 650, "30"),
        box(290, 610, "劲荐"), box(940, 760, "72.2"),
        box(350, 700, "最大外功攻击"), box(920, 830, "114.1"),
        box(340, 770, "转会心率荐"), box(890, 900, "12.7%"),
        box(250, 860, "对首领单位增伤"), box(880, 950, "4.8%"),
        box(250, 950, "劲荐"), box(880, 1030, "72.2"),
        box(410, 1060, "无名剑法·蓄力技增伤荐"), box(870, 1150, "8.9%"),
        box(260, 1150, "易相套装4/4"),
        box(130, 1260, "耐久度"), box(140, 1330, "穿戴等级"),
        box(740, 1390, "001/86"), box(740, 1450, "100级"),
    ]
    piece = parse_ocr_result(ocr, source="t.jpg")
    got = {(a.name, a.value, a.unit) for a in piece.affixes}
    assert ("劲", 72.2, "flat") in got
    assert ("最大外功攻击", 114.1, "flat") in got
    assert ("会心率", 12.7, "percent") in got            # 转会心率
    assert ("对首领单位增伤", 4.8, "percent") in got
    assert ("无名剑法·蓄力技增伤", 8.9, "percent") in got
    assert piece.craft_score == 1476
    assert piece.set_name == "易相套装4/4"
    assert all("气血" not in a.name and "外功防御" not in a.name for a in piece.affixes)


def test_parse_type_constraint_prevents_steal():
    """百分比数值不允许配到固定数值类词条（防错位的关键约束）。"""
    ocr = [
        box(300, 100, "劲"), box(900, 160, "6.6%"),      # 百分比值在“劲”下方
        box(300, 220, "会心率"), box(900, 280, "6.6%"),
    ]
    piece = parse_ocr_result(ocr)
    got = {(a.name, a.value, a.unit) for a in piece.affixes}
    # 劲(flat) 不能吃 6.6%：会心率行配对后，第一行 6.6% 只能被跳过或落到未知名上
    assert ("会心率", 6.6, "percent") in got


def test_parse_marks_converted_and_unknown():
    ocr = [
        box(300, 100, "转会心率"), box(900, 200, "6.6%"),
        box(300, 220, "神秘词条"), box(900, 320, "15.5"),
    ]
    piece = parse_ocr_result(ocr)
    assert piece.affixes[0].converted is True
    assert piece.affixes[0].name == "会心率"
    assert piece.affixes[1].known is False  # 未收录词条保留原文


def test_parse_base_stats_repair_missing_defense():
    """OCR 漏掉外功防御数值时，气血量级的数值应回正到气血最大值行。"""
    ocr = [
        box(310, 320, "装备等阶"),
        box(310, 400, "气血最大值"), box(920, 470, "9723"),   # 外防数值 34 被 OCR 漏掉
        box(280, 460, "外功防御"),
        box(340, 580, "会心率荐"), box(890, 650, "6.6%"),
    ]
    piece = parse_ocr_result(ocr)
    bs = {b.name: b.value for b in piece.base_stats}
    assert bs.get("气血最大值") == 9723.0
    assert "外功防御" not in bs or bs["外功防御"] < 500
    assert any(a.name == "会心率" for a in piece.affixes)


def test_parse_base_stats_swap_when_defense_bigger():
    ocr = [
        box(310, 400, "气血最大值"), box(920, 470, "34"),
        box(280, 460, "外功防御"), box(920, 530, "9723"),
        box(340, 640, "会心率"), box(890, 710, "6.6%"),
    ]
    piece = parse_ocr_result(ocr)
    bs = {b.name: b.value for b in piece.base_stats}
    assert bs.get("气血最大值") == 9723.0
    assert bs.get("外功防御") == 34.0


# ---------------- build import ----------------
@pytest.fixture()
def fake_calc(tmp_path):
    """构造一个与社区计算器同构的最小输入区 xlsx（含第 1 行列头）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "期望"
    ws.cell(row=1, column=2, value="最小攻击")
    ws.cell(row=1, column=3, value="最大攻击")
    ws.cell(row=1, column=4, value="穿透")
    ws.cell(row=1, column=5, value="伤害加成")
    rows = [
        ("外功", 4000, 3100, 60, 0.02),
        ("鸣金", None), ("裂石", None), ("牵丝", None),
        ("破竹", 492, 1260, 30, 0.15),
        ("精准率", 0.99, "实际属性"),
        ("会心率", 0.78),
        ("会意率", 0.09),
        ("直接会心率", 0.13),
        ("直接会意率", 0),
        ("会心伤害加成", 0.54),
        ("会意伤害加成", 0.35),
        ("全武增", 0.0852),
        ("首领增", 0.0887),
        ("拳甲增", 0.0852),
        ("双刀增", 0.0852),
        ("单体奇术增", 0),
        ("群体奇术增", 0),
        ("蓄力技定音", 0.32),
        ("固伤加成", 0.5),
        ("食物加成", 200),
    ]
    for i, r in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=r[0])
        for j, v in enumerate(r[1:], start=2):
            if v is not None:
                ws.cell(row=i, column=j, value=v)
    path = tmp_path / "破竹樽110阶竞速轴属性毕业率进阶计算器2.0.xlsx"
    wb.save(path)
    return path


def test_build_name_from_filename():
    assert _build_name("破竹樽110阶竞速轴属性毕业率进阶计算器2.0.xlsx") == "破竹樽"
    assert _build_name("牵丝玉110阶毕业率.xlsx") == "牵丝玉"
    assert _build_name("破竹鸢2.7.xlsx") == "破竹鸢"  # 简短文件名也兼容


def test_extract_relevant_stats(fake_calc):
    info = extract_relevant_stats(fake_calc)
    assert "最大外功攻击" in info["affixes"]       # 最大攻击
    assert "外功穿透" in info["affixes"]           # 穿透
    assert "会心率" in info["affixes"] and "会意率" in info["affixes"]
    assert "对首领单位增伤" in info["affixes"]     # 首领增
    assert "蓄力技增伤" in info["affixes"]         # 蓄力技定音 -> 特技
    assert "会心伤害" in info["affixes"]


def test_import_build_writes_json(fake_calc, tmp_path):
    out = import_build(fake_calc, out_dir=tmp_path / "builds")
    data = json.loads(out.read_text(encoding="utf-8"))
    assert data["name"] == "破竹樽"
    assert data["affix_weights"]["会心率"] == 1.0


# ---------------- scoring ----------------
MAX_TABLE = {"会心率": 6.6, "劲": 116.3, "无名剑法·蓄力技增伤": 8.9, "最大外功攻击": 114.1}


def _piece_with(affixes):
    from yanyun_gradecalc.ocr import Affix, Piece

    return Piece(name="测试", affixes=[Affix(n, n, v, u) for n, v, u in affixes])


def test_affix_roll():
    from yanyun_gradecalc.ocr import Affix

    assert affix_roll(Affix("会心率", "会心率", 6.6, "percent"), MAX_TABLE) == (100.0, False)
    assert affix_roll(Affix("会心率", "会心率", 3.3, "percent"), MAX_TABLE) == (50.0, False)
    assert affix_roll(Affix("神秘词条", "神秘词条", 1.0, "flat"), MAX_TABLE) == (None, False)
    # 数值远超满值 -> 疑似基础属性行，排除
    assert affix_roll(Affix("最大外功攻击", "最大外功攻击", 199.0, "flat"), MAX_TABLE) == (None, True)


def test_weight_for_special_affix():
    from yanyun_gradecalc.scoring import weight_for

    weights = {"蓄力技增伤": 2.0, "会心率": 1.0}
    assert weight_for("会心率", weights) == 1.0
    assert weight_for("无名剑法·蓄力技增伤", weights) == 2.0  # 特技后缀兜底
    assert weight_for("神秘词条", weights) == 0.0


def test_score_piece_weights_ignore_irrelevant():
    build = {"name": "破竹樽", "affix_weights": {"会心率": 1.0, "劲": 1.0}}
    piece = _piece_with([("会心率", 6.6, "percent"), ("劲", 58.15, "flat"),
                         ("神秘词条", 9.9, "flat")])
    sc = score_piece(piece, build, MAX_TABLE)
    # 神秘词条不被流派看重，不进分母：100% 会心 + 50% 劲
    assert sc.score == 75.0
    assert sc.grade == "可用"
    assert sum(1 for d in sc.details if not d.relevant) == 1


def test_score_piece_missing_max_is_excluded():
    build = {"name": "B", "affix_weights": {"会心率": 1.0, "最大外功攻击": 2.0}}
    piece = _piece_with([("会心率", 6.6, "percent")])
    sc = score_piece(piece, build, MAX_TABLE)  # 无最大外功攻击词条
    assert sc.score == 100.0


def test_compare_builds_ranking():
    builds = {
        "看会心": {"name": "看会心", "affix_weights": {"会心率": 1.0}},
        "看劲": {"name": "看劲", "affix_weights": {"劲": 1.0}},
    }
    piece = _piece_with([("会心率", 6.6, "percent"), ("劲", 11.63, "flat")])
    ranked = compare_builds(piece, builds, MAX_TABLE)
    assert ranked[0].build == "看会心"  # 会心满卷 100 vs 劲 10


def test_grade_thresholds():
    assert grade_of(95) == "毕业"
    assert grade_of(85) == "准毕业"
    assert grade_of(75) == "可用"
    assert grade_of(50) == "过渡"


# ---------------- repo data ----------------
def test_repo_max_table_loads():
    # 满值表由用户在游戏内持续校对，测试只做结构性断言，不锁定具体数值
    t = load_max_table()
    assert t["最大外功攻击"] > 0
    assert t["会意率"] > 0 and t["会心率"] > 0
    assert t["外功穿透"] > 0
