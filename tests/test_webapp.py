"""Web 服务测试：全部离线（OCR 注入 Fake 实现，不加载真实模型）。pytest -q"""

import io
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "webapp"))

import server  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("ADMIN_TOKEN", "test-token")
    monkeypatch.setattr(server, "ROOT", tmp_path)
    (tmp_path / "builds").mkdir(exist_ok=True)
    (tmp_path / "data").mkdir(exist_ok=True)
    (tmp_path / "data" / "affix_max.json").write_text(
        '{"max": {"会心率": 6.6}}', encoding="utf-8")
    # 最小 builds 文件
    import json

    (tmp_path / "builds" / "测试流派.json").write_text(
        json.dumps({"name": "测试流派", "affix_weights": {"会心率": 1.0}},
                   ensure_ascii=False), encoding="utf-8")
    with TestClient(server.app) as c:
        yield c


def test_health_and_index(client):
    assert client.get("/api/health").json() == {"status": "ok"}
    html = client.get("/").text
    assert "毕业度计算器" in html


def test_meta(client):
    m = client.get("/api/meta").json()
    assert "测试流派" in m["builds"]
    assert m["maxes"]["会心率"] == 6.6
    assert "会心率" in m["known_stats"]


FAKE_PIECE = {
    "name": "承音", "slot": "", "tier": "110阶", "set_name": "易相套装4/4",
    "craft_score": 1476, "source": "photo.jpg", "rotated": 0,
    "affixes": [
        {"name": "劲", "raw_name": "劲", "value": 72.2, "unit": "flat",
         "converted": False, "recommended": True, "known": True},
        {"name": "会心率", "raw_name": "转会心率", "value": 12.7, "unit": "percent",
         "converted": True, "recommended": True, "known": True},
    ],
}


def test_ocr_endpoint_with_injected_fake(client):
    server.set_ocr_fn(lambda raw: dict(FAKE_PIECE))
    resp = client.post("/api/ocr", files={
        "image": ("photo.jpg", b"\x89PNG fake bytes", "image/png")})
    assert resp.status_code == 200
    data = resp.json()
    assert data["name"] == "承音"
    assert len(data["affixes"]) == 2
    assert data["affixes"][0]["value"] == 72.2


def test_ocr_invalid_image(client):
    server.set_ocr_fn(lambda raw: (_ for _ in ()).throw(ValueError("无法解码图片")))
    resp = client.post("/api/ocr", files={
        "image": ("x.jpg", b"not-an-image", "image/jpeg")})
    assert resp.status_code == 400


def test_admin_import_requires_token(client, monkeypatch):
    files = {"files": ("calc.xlsx", b"fake", "application/octet-stream")}
    monkeypatch.delenv("ADMIN_TOKEN", raising=False)
    assert client.post("/api/admin/import-builds", files=files).status_code == 503
    monkeypatch.setenv("ADMIN_TOKEN", "test-token")
    assert client.post("/api/admin/import-builds", files=files,
                       headers={"X-Admin-Token": "wrong"}).status_code == 401


def test_admin_import_real_xlsx(client, tmp_path):
    """用最小计算器结构验证导入端到端。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "期望"
    ws.cell(row=1, column=2, value="最小攻击")
    ws.cell(row=1, column=3, value="最大攻击")
    rows = [("外功", 4000, 3100, 60, 0.02), ("会心率", 0.78), ("首领增", 0.0887),
            ("食物加成", 200)]
    for i, r in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=r[0])
        for j, v in enumerate(r[1:], start=2):
            if v is not None:
                ws.cell(row=i, column=j, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    resp = client.post(
        "/api/admin/import-builds?x_admin_token=test-token",
        files={"files": ("新流派110阶计算器.xlsx", buf.getvalue(),
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    assert "新流派" in resp.json()["imported"]
    # 导入后 meta 立即可见
    assert "新流派" in client.get("/api/meta").json()["builds"]
